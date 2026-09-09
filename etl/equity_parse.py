"""
Phase 2, step 2 (spec section 15.8 step 2).

Normalizes the raw NSTA field equity workbook (as downloaded by
etl/equity_fetch.py) into a structured interval table: one row per
(field, company, validity interval), with columns field name, company
name, interest percentage, start date, and end date.

This module does NOT interpret what the rows mean, does NOT join against
production, and does NOT resolve sentinel start dates or field/company
name aliases. It only normalizes and flags:

- open-ended intervals   -> end_date = None (never a sentinel date)
- sentinel start dates   -> start_is_sentinel = True, start_date passed
                             through unchanged (interpretation deferred)
- zero-interest rows     -> is_zero_interest = True, row retained
                             (never dropped, meaning never assumed)

Run: python etl/equity_parse.py [path-to-workbook.xlsx]
"""

from __future__ import annotations

import json
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

DEFAULT_WORKBOOK_PATH = Path(__file__).parent / ".cache" / "equity_workbook.xlsx"
SCHEMA_REPORT_PATH = Path(__file__).parent / "equity_schema_report.md"
INTERVALS_CACHE_PATH = Path(__file__).parent / ".cache" / "equity_intervals.json"

EXPECTED_SHEET_NAME = "Report 1"

# Full set of columns observed live (2026-09-09). Only the first five are
# used by this milestone's normalized interval table (spec's explicit
# minimum); the rest are recorded in the schema report but not carried
# into the parsed output - see etl/equity_schema_report.md "Columns not
# carried into the normalized table".
COL_FIELD_NAME = "Field Name"
COL_COMPANY = "Organisation Name"
COL_INTEREST = "Percentage Holding"
COL_START = "Start Date"
COL_END = "End Date"
EXPECTED_COLUMNS = [
    COL_FIELD_NAME,
    "On Offshore",
    "Median Line Flag",
    "Status",
    COL_COMPANY,
    COL_INTEREST,
    "Operator Flag",
    COL_START,
    COL_END,
    "Equity Share Time Period",
]

# A start date before this year is treated as a placeholder ("unknown /
# since inception"), not a real date - per spec section 15.3, based on the
# archived 2014-2020 workbook. No such rows exist in the live workbook as
# of 2026-09-09 (min Start Date = 1968-08-01; see schema report). This
# still guards against silently mis-parsing one if NSTA reintroduces the
# convention, or if an older archive is ever fed through this parser.
SENTINEL_YEAR_THRESHOLD = 1960


class EquityParseError(RuntimeError):
    """Raised when the workbook cannot be parsed into the expected interval
    structure. Message must say exactly why."""


def _to_iso_date(value, row_num: int, column: str) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raise EquityParseError(
        f"Row {row_num}, column {column!r}: expected a date or blank, got "
        f"{value!r} (type {type(value).__name__})."
    )


def load_workbook_sheet(path: Path):
    if not path.exists():
        raise EquityParseError(
            f"Workbook not found at {path}. Run etl/equity_fetch.py first."
        )
    wb = openpyxl.load_workbook(path, data_only=True)
    if len(wb.sheetnames) != 1:
        raise EquityParseError(
            f"Workbook has {len(wb.sheetnames)} sheet(s) {wb.sheetnames}, "
            "expected exactly 1. Structure may have changed - refusing to "
            "guess which sheet holds the data."
        )
    if wb.sheetnames[0] != EXPECTED_SHEET_NAME:
        raise EquityParseError(
            f"Workbook's single sheet is named {wb.sheetnames[0]!r}, "
            f"expected {EXPECTED_SHEET_NAME!r}."
        )
    return wb[EXPECTED_SHEET_NAME]


def read_header(ws) -> dict:
    """Map expected column name -> 0-based index, by reading row 1. Fails
    loudly if any expected column is missing (deterministic matching only,
    no fuzzy/positional guessing)."""
    header_row = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    header_row = [h.strip() if isinstance(h, str) else h for h in header_row]

    missing = [c for c in EXPECTED_COLUMNS if c not in header_row]
    if missing:
        raise EquityParseError(
            f"Workbook header is missing expected column(s) {missing}. "
            f"Actual header: {header_row}."
        )

    extra = [h for h in header_row if h is not None and h not in EXPECTED_COLUMNS]
    if extra:
        print(f"NOTE: workbook has extra column(s) not modelled here: {extra}")

    return {name: header_row.index(name) for name in EXPECTED_COLUMNS}


def parse_rows(ws, col_index: dict) -> list[dict]:
    rows = []
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None for v in row):
            continue  # trailing blank row

        field_name = row[col_index[COL_FIELD_NAME]]
        company_name = row[col_index[COL_COMPANY]]
        interest_pct = row[col_index[COL_INTEREST]]
        start_raw = row[col_index[COL_START]]
        end_raw = row[col_index[COL_END]]

        if not field_name:
            raise EquityParseError(f"Row {row_num}: missing {COL_FIELD_NAME!r}.")
        if not company_name:
            raise EquityParseError(f"Row {row_num}: missing {COL_COMPANY!r}.")
        if interest_pct is None or isinstance(interest_pct, bool) or not isinstance(interest_pct, (int, float)):
            raise EquityParseError(
                f"Row {row_num}: {COL_INTEREST!r} is {interest_pct!r}, "
                "expected a number. Zero is valid and must be retained - "
                "only a missing/non-numeric value is an error."
            )
        if interest_pct < 0 or interest_pct > 100:
            raise EquityParseError(
                f"Row {row_num}: {COL_INTEREST!r} {interest_pct!r} is out "
                "of the expected [0, 100] range."
            )
        if start_raw is None:
            raise EquityParseError(
                f"Row {row_num}: {COL_START!r} is missing - every interval "
                "must have a start date; only End Date may be open-ended."
            )

        start_date = _to_iso_date(start_raw, row_num, COL_START)
        end_date = _to_iso_date(end_raw, row_num, COL_END)

        if end_date is not None and end_date < start_date:
            raise EquityParseError(
                f"Row {row_num}: End Date {end_date} precedes Start Date "
                f"{start_date}."
            )

        start_is_sentinel = isinstance(start_raw, (date, datetime)) and start_raw.year < SENTINEL_YEAR_THRESHOLD

        rows.append(
            {
                "field_name": field_name,
                "company_name": company_name,
                "interest_pct": float(interest_pct),
                "start_date": start_date,
                "end_date": end_date,
                "start_is_sentinel": start_is_sentinel,
                "is_zero_interest": float(interest_pct) == 0.0,
            }
        )
    return rows


def parse_equity_workbook(path: Path = DEFAULT_WORKBOOK_PATH) -> list[dict]:
    ws = load_workbook_sheet(path)
    col_index = read_header(ws)
    return parse_rows(ws, col_index)


def build_schema_report(path: Path, ws, rows: list[dict]) -> str:
    header_row = list(next(ws.iter_rows(min_row=1, max_row=1, values_only=True)))
    open_ended = [r for r in rows if r["end_date"] is None]
    sentinel = [r for r in rows if r["start_is_sentinel"]]
    zero_interest = [r for r in rows if r["is_zero_interest"]]
    start_dates = [r["start_date"] for r in rows]
    end_dates = [r["end_date"] for r in rows if r["end_date"] is not None]
    fields = sorted({r["field_name"] for r in rows})
    companies = sorted({r["company_name"] for r in rows})

    lines = [
        "# NSTA field equity workbook - schema report",
        "",
        f"- Workbook file: `{path.name}`",
        f"- Sheet names: {ws.parent.sheetnames if hasattr(ws, 'parent') else [ws.title]}",
        f"- Column names (in order): {header_row}",
        f"- Data row count: {len(rows)}",
        f"- Distinct field names: {len(fields)}",
        f"- Distinct company (organisation) names: {len(companies)}",
        f"- Start date range: {min(start_dates)} to {max(start_dates)}" if start_dates else "- Start date range: (no rows)",
        f"- End date range (closed intervals only): {min(end_dates)} to {max(end_dates)}" if end_dates else "- End date range: (no closed intervals)",
        f"- Open-ended rows (End Date blank): {len(open_ended)}",
        f"- Sentinel start-date rows (< {SENTINEL_YEAR_THRESHOLD}): {len(sentinel)}",
        f"- Zero-interest rows (Percentage Holding = 0): {len(zero_interest)}",
        "",
        "## Columns not carried into the normalized table",
        "",
        "The live workbook has 5 columns beyond the milestone's required "
        "set (field name, company name, interest %, start date, end date): "
        "`On Offshore`, `Median Line Flag`, `Status`, `Operator Flag`, "
        "`Equity Share Time Period`. These are present in the source and "
        "recorded here for visibility, but are not part of this "
        "milestone's normalized output.",
        "",
        "## Deviations from spec section 15 assumptions",
        "",
        "1. **Section 15.1 assumes two separate NSTA datasets** - a "
        "time-series \"Field Equity Shares\" dataset (the correct primary "
        "source) and a present-day-only \"Field Partners\" snapshot "
        "(forbidden for time series, cross-check only). Live discovery "
        "found only **one** dataset in the NSTA ArcGIS org "
        "(`OZMfUznmLTnWccBc`): an item titled **\"Field Partners\"** "
        "(id `40fb75005dca48e886891350da9dedd8`). Searching that org for "
        "the exact phrase \"field equity shares\" returns zero results. "
        "The NSTA Fields page's own link, labelled \"Current and "
        "historical field equity shares\", resolves to this same item. "
        "Its actual content is full interval history (Start Date back to "
        "1968-08-01, 277 distinct \"Equity Share Time Period\" labels such "
        "as \"Previous-2005 to 2020\" and \"Current\"), not a present-day "
        "snapshot. Section 15.1's snapshot description does not match "
        "what this item currently contains. This is the only live source "
        "available and it satisfies section 15.1's structural "
        "requirements (interest %, start date, end date, records to the "
        "late 1960s) despite not being named \"Field Equity Shares\".",
        "",
        "2. **Section 15.2 assumes a single-page scrape to a direct file "
        "link.** In reality, resolution is two-hop: NSTA page -> ArcGIS "
        "Hub search page (client-rendered SPA, no server-side link to "
        "scrape) -> ArcGIS Online search API scoped to org "
        "`OZMfUznmLTnWccBc` -> resolved item -> direct `.xlsx` URL on "
        "Azure Blob Storage (`datanstauthority.blob.core.windows.net`, "
        "not `nstauthority.co.uk`). `etl/equity_fetch.py` implements this "
        "two-hop resolution.",
        "",
        f"3. **No sentinel (`1900-01-01`-style) start dates found.** "
        f"Section 15.2/15.3/15.4 describe such rows based on the archived "
        f"2014-2020 file. Zero rows in the live {len(rows)}-row dataset "
        f"have a start date before {SENTINEL_YEAR_THRESHOLD} (earliest "
        f"observed: {min(start_dates) if start_dates else 'n/a'}). The "
        "parser still detects them defensively; it currently never fires.",
        "",
        "4. **4 rows have a future Start Date** relative to today "
        "(2026-09-09): two starting 2030-05-01 (ALVHEIM / AKER BP ASA; "
        "STATFJORD(CROSS BORDER) / EQUINOR UK LIMITED) and two starting "
        "2050-01-04 (MURLACH [pt of MARNOCK-SKUA] / BP EXPLORATION "
        "OPERATING COMPANY LIMITED and NEO ENERGY (ZNS) LIMITED). Not "
        "addressed by any rule in section 15.4. Not resolved in this "
        "milestone.",
        "",
        f"5. **Zero-interest rows confirmed to exist**: {len(zero_interest)} "
        f"of {len(rows)} rows ({len(zero_interest) / len(rows):.1%}). "
        "Section 15.4 asks that their meaning be investigated before "
        "deciding how to treat them. Not investigated here per this "
        "milestone's explicit scope; rows are retained, unfiltered, with "
        "`is_zero_interest=true`. Notably, some zero-interest rows have "
        "`Operator Flag = 'Y'` (e.g. AFFLECK / NEO NEXT+ ENERGY PUK "
        "LIMITED, current row) - an operator with no recorded equity "
        "interest, reinforcing that these rows should not be assumed to "
        "mean \"relinquished\" or \"nulled\" without further investigation.",
        "",
        "6. **At least one zero-duration interval exists** (Start Date == "
        "End Date, e.g. ALISON [CENTRICA] / SPIRIT NORTH SEA GAS LIMITED, "
        "both dated 2016-05-17). Section 15.4's proposed rule E7 "
        "(`start_date < end_date`, strict) would fail this row. Not fixed "
        "here - flagged for whoever implements E7.",
        "",
        f"7. **Scale for later name reconciliation (not attempted here)**: "
        f"{len(fields)} distinct field names and {len(companies)} distinct "
        "organisation names in the live workbook, versus 552 fields in "
        "PPRS full history (per README). Field/company matching is "
        "explicitly out of scope for this milestone (spec section 15.8 "
        "step 3).",
        "",
    ]
    return "\n".join(lines)


def main() -> int:
    argv_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_WORKBOOK_PATH
    try:
        ws = load_workbook_sheet(argv_path)
        col_index = read_header(ws)
        rows = parse_rows(ws, col_index)

        if not rows:
            raise EquityParseError(f"Parsed zero data rows from {argv_path}.")

        report = build_schema_report(argv_path, ws, rows)
        SCHEMA_REPORT_PATH.write_text(report)
        print(report)
        print(f"\nSchema report written to {SCHEMA_REPORT_PATH}")

        INTERVALS_CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
        INTERVALS_CACHE_PATH.write_text(json.dumps(rows, indent=2, sort_keys=True))
        print(f"Normalized interval table ({len(rows)} rows) written to {INTERVALS_CACHE_PATH}")

        return 0
    except EquityParseError as e:
        print(f"\nEQUITY PARSE FAILED: {e}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
